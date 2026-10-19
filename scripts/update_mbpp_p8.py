# -*- coding: utf-8 -*-
"""将报告表 4-3 的 MBPP pass@8 从占位更新为真实结果（std_mbpp_s8.json 完成后运行）。"""
import json
import sys
from pathlib import Path

import docx

SRC = Path(r"F:\自研大模型\docs\自研106M大模型数据侧能力迭代研究报告.docx")
MB8 = Path(r"F:\自研大模型\logs\std_mbpp_s8.json")

if not MB8.exists():
    print("std_mbpp_s8.json 尚未落盘，终止（占位保持不变）")
    sys.exit(1)

d = json.load(open(MB8, encoding="utf-8"))
# std JSON 无 done 字段：完成才落盘且 n=257 视为完成
if d["meta"].get("n", 0) != 257:
    print(f"n={d['meta'].get('n')} != 257，非全量，终止")
    sys.exit(1)
m = d["meta"]
s = d["summary"]
val = f"{m['n_pass']}/{m['n']} ({m['pass_rate']*100:.0f}%)"
print("MBPP p8 真实值:", val, "| summary:", s)

doc = docx.Document(str(SRC))
hit = 0
for tb in doc.tables:
    for row in tb.rows:
        cells = [c.text.strip() for c in row.cells]
        if any("MBPP" in c for c in cells):
            for ci, c in enumerate(row.cells):
                if "待补" in c.text or "后台运行" in c.text:
                    for p in c.paragraphs:
                        for r in p.runs:
                            if "待补" in r.text or "后台运行" in r.text:
                                r.text = val
                                hit += 1
                            else:
                                r.text = ""
if hit == 0:
    print("未找到占位单元格（可能已更新）")
else:
    doc.save(str(SRC))
    print(f"已更新 {hit} 个 run -> 表 4-3 MBPP pass@8 = {val}")
