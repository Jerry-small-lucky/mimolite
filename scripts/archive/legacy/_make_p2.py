# -*- coding: utf-8 -*-
"""生成 2 页研究提案 docx（预算守恒与函数名竞争）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"F:\自研大模型\docs\研究提案-预算守恒与函数名竞争.docx"

doc = Document()
# 页面 A4 2.5cm
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)

def set_font(run, cn, en="Times New Roman", size=12, bold=False, color=(0, 0, 0)):
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run(text), "黑体", size=18, bold=True)
    return p

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), "黑体", size=14, bold=True)
    return p

def body(text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.3
    pf.space_before = pt0 = Pt(0)
    pf.space_after = Pt(2)
    if indent:
        pf.first_line_indent = Pt(24)
    set_font(p.add_run(text), "宋体", size=12)
    return p

def ref(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.1
    pf.space_after = Pt(2)
    pf.left_indent = Pt(0)
    set_font(p.add_run(text), "宋体", size=10.5)
    return p

title("预算守恒与函数名竞争：极小模型合成数据工程的样本—函数分配")

h1("一、问题与动机")
body("Self-Instruct [1] 与 phi-1 [2] 等研究表明，合成数据可以显著降低指令微调与代码生成对人工标注的依赖；"
     "STaR [3] 进一步证明模型可通过自举提升推理能力。但三条主流路线都依赖教师模型（GPT-4/GPT-3.5）或自举信号，"
     "且模型规模通常在 1B 以上。工业界的大规模后训练（如 MiMo-V2.5 [4]）则依赖 48T token 与百亿级模型。")
body("与之相对，本研究占据一个几乎无人系统研究的约束组合：106M 参数、无教师蒸馏、纯规则化合成数据、单卡 8GB 显存。"
     "在 8 版同口径实验中（表 1），数据多样性（v6）与按评测频次加权（v7）将 pass@1 从 27% 提升至 30%；"
     "但两次负优化实验给出了更重要的信息：v9 删除干扰变体并将 3 个目标函数权重提至 32，pass@1 从 29% 暴跌至 20%；"
     "v10 扩大总量至 720 并将目标函数配额提至 40，全局仍回落到 28%/51%。")
body("这两个负结果说明：固定合成预算下“强化 A 必削弱 B”（预算守恒），且被强化函数若存在能力瓶颈，"
     "加样本不会转化为通过率（函数名竞争）。本研究将这些现象形式化为可检验命题，并给出预算最低的验证方案。")

h1("二、可检验命题")
body("H1（预算守恒律）：固定合成预算 B 下，对函数集合 F 的样本分配存在帕累托最优。"
     "预测：“按评测频次加权”优于“均匀分配”与“定向强化”；全局 pass@1/pass@8 随高频权重 w 呈单峰，"
     "当前最优 w*=16（v8 配置）应位于峰值附近。")
body("H2（函数名竞争律）：语义同质、函数名不同的函数对（如 count_chars 与 count_characters），"
     "小模型只能学好样本占优的名字（保大弃小），通过率随两名字的训练占比存在相变阈值。"
     "预测：占比低于阈值的名字 pass@1≈0，跨越阈值后快速恢复。")
body("H3（多采样恢复，观测性）：pass@8 能恢复部分 pass@1 损失。已有证据：51 个评测函数中 13 个 pass@1=0 而 pass@8>0"
     "（process_string 4/4、fib 4/4 等，表 6），说明能力已灌入权重，pass@1 受单采样确定性低估。")

h1("三、实验方案")
body("全部实验复用现有管线，保证无泄漏与同口径：synth_code_tpl.py（--n/--seed/--out）→ build_train_v2.py → "
     "prepare_sft.py --min 1 → sft.py --epochs 10（330 步，单轮约 40 分钟）→ eval_code.py（p1/p8，固定采样参数）。")
body("实验 1（H1）：固定 B=600 条，扫描高频权重 w∈{6,8,16,24,32} 共 5 轮训练，测全局 pass@1/pass@8 曲线，定位峰值 w*。")
body("实验 2（H2）：固定同质函数对总预算，扫描名字占比 p∈{0.1,0.3,0.5,0.7,0.9} 共 5 轮训练，"
     "测两函数各自通过率，定位相变区间。")
body("实验 3（H3）：复用 v8 既有档案（logs/eval_code_sftv8_p8.json）作观测，无需新训练。"
     "总预算：约 10 轮训练 ≈ 7 小时 GPU 时间（RTX 5070 8GB，单卡串行）。")

h1("四、预期贡献")
body("其一，为“极小模型 + 无蒸馏 + 纯数据工程”提供第一性设计准则：预算守恒律（分配优于总量）"
     "与函数名竞争律（函数名是第一优先级特征）。其二，产出可复现的 SFT 数据配置指南（权重扫描曲线与占比相变数据），"
     "对低资源代码模型训练有直接实践价值。其三，填补“合成数据预算分配”的研究空白，"
     "为后续 GRPO 自举与模型规模升级（已具备管线）提供可对接的前置结论。")

h1("参考文献")
ref("[1] Wang Y, Kordi Y, Mishra S, et al. Self-Instruct: Aligning Language Models with Self-Generated Instructions[J]. "
    "arXiv preprint arXiv:2212.10560, 2022.")
ref("[2] Gunasekar S, Zhang Y, Aneja J, et al. Textbooks Are All You Need[J]. arXiv preprint arXiv:2306.11644, 2023.")
ref("[3] Zelikman E, Wu Y, Mu J, et al. STaR: Bootstrapping Reasoning With Reasoning[C]//NeurIPS 2022. arXiv:2203.14465.")
ref("[4] Xiaomi. MiMo-V2.5 技术主页[EB/OL]. (2026-04-22). https://mimo.xiaomi.com/mimo-v2-5/.")

doc.save(OUT)
print("saved", OUT)
