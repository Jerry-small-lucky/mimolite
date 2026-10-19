# -*- coding: utf-8 -*-
"""修复表编号：复现命令链=表7，GRPO矩阵=表8"""
from docx import Document

TARGET = r"F:\自研大模型\docs\自研106M大模型数据侧能力迭代研究报告.docx"
doc = Document(TARGET)

fixed = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith("表") and "GRPO" in t:
        for r in p.runs:
            r.text = r.text.replace("表 7", "表 8")
        fixed.append(("GRPO→表8", p.text[:40]))
    elif t.startswith("表") and "复现命令链" in t:
        for r in p.runs:
            r.text = r.text.replace("表 8", "表 7")
        fixed.append(("复现→表7", p.text[:40]))

# 检查正文引用残留
residual = [p.text[:60] for p in doc.paragraphs if "如表 7 所示" in p.text]
doc.save(TARGET)
print(fixed)
print("残留'如表7所示':", residual if residual else "无")
