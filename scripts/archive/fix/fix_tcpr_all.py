# -*- coding: utf-8 -*-
"""修复全部表格所有单元格 tcPr 子节点为 OOXML schema 顺序"""
from docx import Document
from docx.oxml.ns import qn

TARGET = r"F:\自研大模型\docs\自研106M大模型数据侧能力迭代研究报告.docx"
doc = Document(TARGET)

# OOXML CT_TcPr 子元素顺序
ORDER = [qn("w:tcW"), qn("w:gridSpan"), qn("w:hMerge"), qn("w:vMerge"), qn("w:tcBorders"),
         qn("w:shd"), qn("w:noWrap"), qn("w:tcMar"), qn("w:textDirection"), qn("w:tcFitText"),
         qn("w:vAlign"), qn("w:hideMark")]

def fix_cell(cell):
    tcPr = cell._tc.find(qn("w:tcPr"))
    if tcPr is None:
        return 0
    nodes = list(tcPr)
    nodes.sort(key=lambda n: ORDER.index(n.tag) if n.tag in ORDER else 99)
    moved = sum(1 for a, b in zip(nodes, list(tcPr)) if a is not b)
    for n in nodes:
        tcPr.remove(n)
    for n in nodes:
        tcPr.append(n)
    return moved

total = 0
for ti, tb in enumerate(doc.tables):
    for tr in tb.rows:
        for tc in tr.cells:
            total += fix_cell(tc)
doc.save(TARGET)
print(f"修复完成：全表 {len(doc.tables)} 张，移动节点 {total} 处")
